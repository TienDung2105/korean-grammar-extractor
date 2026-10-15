# -*- coding: utf-8 -*-
"""
Korean Double Negation Web Extractor
=====================================
Web app cho phep upload file JSON va tim cau tieng Han chua cau truc phu dinh kep.
"""

import io
import sys
import os
import json
import re
import zipfile
import tempfile
import shutil
import uuid
import threading
import time
from pathlib import Path
from datetime import datetime
from collections import Counter

from flask import (
    Flask, render_template, request, jsonify, send_file,
    Response, stream_with_context
)

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

app = Flask(__name__)
# Không giới hạn dung lượng upload
app.config['MAX_CONTENT_LENGTH'] = None

@app.errorhandler(413)
def request_entity_too_large(error):
    return jsonify({'error': 'Dung lượng tải lên quá lớn! Đối với thư mục chứa hàng nghìn file JSON, bạn hãy copy đường dẫn thư mục và dán vào ô bên dưới thay vì kéo thả tải lên.'}), 413

@app.errorhandler(500)
def internal_error(error):
    return jsonify({'error': 'Lỗi server: ' + str(error)}), 500

@app.after_request
def add_no_cache_headers(response):
    response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, post-check=0, pre-check=0, max-age=0'
    response.headers['Pragma'] = 'no-cache'
    response.headers['Expires'] = '-1'
    return response

# Store for ongoing jobs
jobs = {}

# ============================================================================
# REGEX PATTERNS - 16 cau truc phu dinh kep (DA TOI UU DO CHINH XAC CAO)
# ============================================================================

# Tap hop toan bo cac chu Hangul mang patchim ㄹ va ㄴ trong Unicode (399 chu moi loai)
RIEUL_CHARS = "".join([chr(c) for c in range(0xac00, 0xd7a4) if (c - 0xac00) % 28 == 8])
NIEUN_CHARS = "".join([chr(c) for c in range(0xac00, 0xd7a4) if (c - 0xac00) % 28 == 4])

# Dong tu/tinh tu co patchim ㄹ (vi du: 갈, 할, 볼) hoac theo sau boi 을 (vi du: 먹을)
V_RIEUL = rf'[가-힣]*(?:[{RIEUL_CHARS}]|[가-힣]+을)'

# Dong tu/tinh tu dang dinh ngu: patchim ㄴ (간, 한, 본) hoac 은/는 (먹은, 가는)
V_MOD = rf'[가-힣]*(?:[{NIEUN_CHARS}]|[가-힣]+(?:은|는))'

# Dong tu/tinh tu tieng Han co ban (1-10 ky tu)
V_WORD = r'[가-힣]{1,10}'

# Cac duoi chia pho bien cua vi tu phu dinh
AN_ENDINGS = r'않[다아고은는을았었겠지며자도서면니습까요야네죠거든래]?'
MOT_ENDINGS = r'(?:못하[다아고는을였겠지며면니습까요야네죠거든래]?|못했[다어고는을겠지며면니습까요야네죠거든래]?|못\s*(?:하[다아고는을였겠지며면니습까요야네죠거든래]?|했[다어고는을겠지며면니습까요야네죠거든래]?|해[다요]?|한))'
MAL_ENDINGS = r'(?:말[다아고자라려면지며니습까요야았겠네죠거든래]?|마[라아요]?)'
EOP_ENDINGS = r'없[다어고는을었겠지며면니습까요야네죠거든래]?'
ANI_ENDINGS = r'아니[다야고라면며는ㄴ란린ㄹ었겠지니습까요네죠거든래]?'

ALL_PATTERNS = [
    {
        'id': '1', 'name': '안 –지 않다', 'desc': '안 V-지 않다',
        'regex': re.compile(r'(?<![가-힣])안\s+' + V_WORD + r'지(?:는|도)?\s*' + AN_ENDINGS, re.UNICODE)
    },
    {
        'id': '2', 'name': '안 –지 못하다', 'desc': '안 V-지 못하다',
        'regex': re.compile(r'(?<![가-힣])안\s+' + V_WORD + r'지(?:는|도)?\s*' + MOT_ENDINGS, re.UNICODE)
    },
    {
        'id': '3', 'name': '못 –지 않다', 'desc': '못 V-지 않다',
        'regex': re.compile(r'(?<![가-힣])(?:못\s+' + V_WORD + r'지|못지)(?:는|도)?\s*' + AN_ENDINGS, re.UNICODE)
    },
    {
        'id': '4', 'name': '안 –지 말-', 'desc': '안 V-지 말다',
        'regex': re.compile(r'(?<![가-힣])안\s+' + V_WORD + r'지(?:는|도)?\s*' + MAL_ENDINGS, re.UNICODE)
    },
    {
        'id': '5', 'name': '없-지 않다', 'desc': '없-지 않다',
        'regex': re.compile(r'(?<![가-힣])[가-힣]*없지(?:는|도)?\s*' + AN_ENDINGS, re.UNICODE)
    },
    {
        'id': '6', 'name': '모르-지 않다', 'desc': '모르-지 않다',
        'regex': re.compile(r'(?<![가-힣])모르지(?:는|도)?\s*' + AN_ENDINGS, re.UNICODE)
    },
    {
        'id': '7', 'name': '안+수/리+없다', 'desc': '안 V-(으)ㄹ 수/리 없다',
        'regex': re.compile(r'(?<![가-힣])안\s+' + V_RIEUL + r'\s*(?:수|리)(?:가|는|도)?\s*' + EOP_ENDINGS, re.UNICODE)
    },
    {
        'id': '8', 'name': '못+수/리+없다', 'desc': '못 V-(으)ㄹ 수/리 없다',
        'regex': re.compile(r'(?<![가-힣])못\s+' + V_RIEUL + r'\s*(?:수|리)(?:가|는|도)?\s*' + EOP_ENDINGS, re.UNICODE)
    },
    {
        'id': '9', 'name': '못+것+아니다', 'desc': '못 V 것 아니다',
        'regex': re.compile(r'(?<![가-힣])못\s+' + V_MOD + r'\s*(?:것|건|게|거)(?:이|은|는|도|가)?\s*' + ANI_ENDINGS, re.UNICODE)
    },
    {
        'id': '10', 'name': '안+것+아니다', 'desc': '안 V 것 아니다',
        'regex': re.compile(r'(?<![가-힣])안\s+' + V_MOD + r'\s*(?:것|건|게|거)(?:이|은|는|도|가)?\s*' + ANI_ENDINGS, re.UNICODE)
    },
    {
        'id': '11', 'name': '없-+것+아니다', 'desc': '없- 것 아니다',
        'regex': re.compile(r'(?<![가-힣])[가-힣]*없는\s*(?:것|건|게|거)(?:이|은|는|도|가)?\s*' + ANI_ENDINGS, re.UNICODE)
    },
    {
        'id': '12', 'name': '모르-+것+아니다', 'desc': '모르- 것 아니다',
        'regex': re.compile(r'(?<![가-힣])모르는\s*(?:것|건|게|거)(?:이|은|는|도|가)?\s*' + ANI_ENDINGS, re.UNICODE)
    },
    {
        'id': '13', 'name': '아니-+것+아니다', 'desc': '아니- 것 아니다',
        'regex': re.compile(r'(?<![가-힣])(?:아닌|아니라는)\s*(?:것|건|게|거)(?:이|은|는|도|가)?\s*' + ANI_ENDINGS, re.UNICODE)
    },
    {
        'id': '14', 'name': '없-+수/리+없다', 'desc': '없- 수/리 없다',
        'regex': re.compile(r'(?<![가-힣])[가-힣]*없을\s*(?:수|리)(?:가|는|도)?\s*' + EOP_ENDINGS, re.UNICODE)
    },
    {
        'id': '15', 'name': '모르-+수/리+없다', 'desc': '모르- 수/리 없다',
        'regex': re.compile(r'(?<![가-힣])모를\s*(?:수|리)(?:가|는|도)?\s*' + EOP_ENDINGS, re.UNICODE)
    },
    {
        'id': '16', 'name': '아니-+수/리+없다', 'desc': '아니- 수/리 없다',
        'regex': re.compile(r'(?<![가-힣])아닐\s*(?:수|리)(?:가|는|도)?\s*' + EOP_ENDINGS, re.UNICODE)
    },
]


# ============================================================================
# PROCESSING FUNCTIONS
# ============================================================================

def extract_year_from_path(filepath):
    path_str = str(filepath)
    year_in_path = re.findall(r'[/\\](20[12]\d)[/\\]', path_str)
    if year_in_path:
        return year_in_path[-1]
    basename = os.path.basename(path_str)
    m = re.search(r'SDRW(\d{2})', basename)
    if m:
        return '20' + m.group(1)
    m = re.search(r'(20[12]\d)', path_str)
    if m:
        return m.group(1)
    return 'N/A'


def extract_year_from_json(data):
    try:
        year = data.get('metadata', {}).get('year', '')
        if year and re.match(r'20\d{2}', str(year)):
            return str(year)
    except:
        pass
    return None


def get_utterances(data):
    utterances = []
    documents = data.get('document', [])
    if not isinstance(documents, list):
        documents = [documents]
    for doc in documents:
        utts = doc.get('utterance', [])
        if not isinstance(utts, list):
            utts = [utts]
        for utt in utts:
            form = utt.get('form', '').strip()
            original_form = utt.get('original_form', '').strip()
            utt_id = utt.get('id', '')
            if form:
                utterances.append({
                    'form': form,
                    'original_form': original_form,
                    'id': utt_id
                })
    return utterances


def search_patterns_in_text(text, patterns):
    matches = []
    for pat in patterns:
        m = pat['regex'].search(text)
        if m:
            matches.append({
                'pattern_id': pat['id'],
                'pattern_name': pat['name'],
                'matched_text': m.group()
            })
    return matches


def process_single_file(filepath, file_content, patterns):
    results = []
    try:
        if file_content is not None:
            data = json.loads(file_content)
        else:
            with open(filepath, 'r', encoding='utf-8') as f:
                data = json.load(f)
    except Exception:
        return results

    year = extract_year_from_json(data) or extract_year_from_path(filepath)
    file_id = data.get('id', os.path.basename(filepath))
    category = data.get('metadata', {}).get('category', '')
    utterances = get_utterances(data)

    for utt in utterances:
        matches = search_patterns_in_text(utt['form'], patterns)
        if not matches and utt['original_form']:
            matches = search_patterns_in_text(utt['original_form'], patterns)
        for match in matches:
            results.append({
                'year': year,
                'file': os.path.basename(filepath),
                'file_id': file_id,
                'category': category,
                'utterance_id': utt['id'],
                'sentence_form': utt['form'],
                'sentence_original': utt['original_form'],
                'pattern_name': match['pattern_name'],
                'matched_text': match['matched_text']
            })
    return results


def save_to_excel(results):
    """Save results to an Excel file in memory, return BytesIO."""
    output = io.BytesIO()

    if not HAS_OPENPYXL:
        # Fallback to CSV
        import csv
        text_output = io.StringIO()
        writer = csv.DictWriter(text_output, fieldnames=[
            'year', 'file', 'utterance_id', 'category',
            'sentence_form', 'sentence_original', 'pattern_name', 'matched_text'
        ])
        writer.writeheader()
        for r in results:
            writer.writerow({k: r[k] for k in writer.fieldnames})
        output = io.BytesIO(text_output.getvalue().encode('utf-8-sig'))
        return output, 'csv'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ket qua"

    headers = [
        ('STT', 6), ('Nam', 8), ('File nguon', 22), ('Ma phat ngon', 28),
        ('The loai', 25), ('Cau tieng Han (chuan hoa)', 55),
        ('Cau tieng Han (goc)', 55), ('Cau truc tim thay', 22), ('Phan khop', 25),
    ]

    header_font = Font(name='Malgun Gothic', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='B4C6E7'),
        right=Side(style='thin', color='B4C6E7'),
        top=Side(style='thin', color='B4C6E7'),
        bottom=Side(style='thin', color='B4C6E7'),
    )

    for col_idx, (header_text, width) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    ws.freeze_panes = 'A2'

    data_font = Font(name='Malgun Gothic', size=10)
    data_font_kr = Font(name='Malgun Gothic', size=11)
    highlight_font = Font(name='Malgun Gothic', size=10, bold=True, color='C00000')
    even_fill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')

    for row_idx, r in enumerate(results, 2):
        row_data = [
            row_idx - 1, r['year'], r['file'], r['utterance_id'],
            r['category'], r['sentence_form'], r['sentence_original'],
            r['pattern_name'], r['matched_text'],
        ]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if col_idx in (6, 7):
                cell.font = data_font_kr
            elif col_idx == 9:
                cell.font = highlight_font
            else:
                cell.font = data_font
            if row_idx % 2 == 0:
                cell.fill = even_fill

    # Stats sheet
    ws2 = wb.create_sheet(title="Thong ke")
    ws2.cell(row=1, column=1, value="Cau truc").font = Font(bold=True)
    ws2.cell(row=1, column=2, value="So cau").font = Font(bold=True)
    ws2.column_dimensions['A'].width = 25
    ws2.column_dimensions['B'].width = 15

    pattern_counts = Counter(r['pattern_name'] for r in results)
    current_row = 2
    for pat_name, cnt in pattern_counts.most_common():
        ws2.cell(row=current_row, column=1, value=pat_name)
        ws2.cell(row=current_row, column=2, value=cnt)
        current_row += 1

    total_row = current_row
    ws2.cell(row=total_row, column=1, value="TONG CONG").font = Font(bold=True)
    ws2.cell(row=total_row, column=2, value=len(results)).font = Font(bold=True)

    wb.save(output)
    output.seek(0)
    return output, 'xlsx'


def build_custom_pattern(raw_text, custom_id):
    """Tu dong tao regex thong minh tu chuoi ngu phap nguoi dung tu nhap."""
    raw = raw_text.strip()
    if not raw:
        return None

    name = raw.strip('-~ ')
    p = raw.strip('~ ')
    if p.startswith('-'):
        p = p[1:]

    # Xu ly cac ky hieu ngu phap pho bien tieng Han
    p = p.replace('(으)ㄹ', f'(?:[{RIEUL_CHARS}]|[가-힣]+을)')
    p = p.replace('(으)ㄴ', f'(?:[{NIEUN_CHARS}]|[가-힣]+은)')
    p = p.replace('(으)', '(?:으)?')

    parts = p.split()
    escaped_parts = []
    for part in parts:
        if RIEUL_CHARS in part or NIEUN_CHARS in part or '(?:' in part:
            escaped_parts.append(part)
        else:
            escaped_parts.append(re.escape(part))
    regex_str = r'\s*'.join(escaped_parts)

    try:
        rx = re.compile(regex_str, re.UNICODE)
    except Exception:
        rx = re.compile(re.escape(name), re.UNICODE)

    return {
        'id': f'custom_{custom_id}',
        'name': name,
        'desc': f'Tự nhập: {name}',
        'regex': rx
    }


# ============================================================================
# BACKGROUND JOB PROCESSOR
# ============================================================================

def process_job(job_id, files_data, folder_path, selected_pattern_ids, custom_patterns_list):
    """Process files in background thread."""
    job = jobs[job_id]
    job['status'] = 'processing'
    job['results'] = []

    # Danh sach pattern tong hop
    active_patterns = []

    if selected_pattern_ids:
        active_patterns.extend([p for p in ALL_PATTERNS if p['id'] in selected_pattern_ids])
    else:
        active_patterns.extend(ALL_PATTERNS)

    # Them cac custom patterns
    if custom_patterns_list:
        active_patterns.extend(custom_patterns_list)

    all_results = []

    # Source 1: uploaded files
    if files_data:
        total = len(files_data)
        job['total_files'] = total
        for idx, (filename, content) in enumerate(files_data):
            job['processed_files'] = idx + 1
            job['current_file'] = filename
            results = process_single_file(filename, content, active_patterns)
            all_results.extend(results)

    # Source 2: folder path
    if folder_path and os.path.isdir(folder_path):
        json_files = sorted(Path(folder_path).rglob('*.json'))
        total = len(json_files)
        job['total_files'] = job.get('total_files', 0) + total
        for idx, fp in enumerate(json_files):
            job['processed_files'] = job.get('processed_files', 0) + 1
            job['current_file'] = fp.name
            results = process_single_file(str(fp), None, active_patterns)
            all_results.extend(results)

    job['results'] = all_results
    job['status'] = 'done'

    # Build stats
    pattern_counts = Counter(r['pattern_name'] for r in all_results)
    year_counts = Counter(r['year'] for r in all_results)
    job['stats'] = {
        'total_sentences': len(all_results),
        'by_pattern': dict(pattern_counts.most_common()),
        'by_year': dict(sorted(year_counts.items())),
    }


# ============================================================================
# ROUTES
# ============================================================================

@app.route('/')
def index():
    return render_template('index.html', patterns=ALL_PATTERNS)


@app.route('/api/process', methods=['POST'])
def api_process():
    """Start processing files."""
    job_id = str(uuid.uuid4())[:8]
    jobs[job_id] = {
        'status': 'starting',
        'total_files': 0,
        'processed_files': 0,
        'current_file': '',
        'results': [],
        'stats': {},
    }

    # Get selected patterns
    selected_ids = request.form.getlist('patterns[]')

    # Get custom patterns
    custom_inputs = request.form.getlist('custom_patterns[]')
    custom_patterns_list = []
    for idx, cp_text in enumerate(custom_inputs, 1):
        built = build_custom_pattern(cp_text, idx)
        if built:
            custom_patterns_list.append(built)

    # Collect uploaded files
    files_data = []
    uploaded = request.files.getlist('files')
    for f in uploaded:
        if not f.filename:
            continue
        content = f.read()

        # Handle ZIP files
        if f.filename.lower().endswith('.zip'):
            try:
                zf = zipfile.ZipFile(io.BytesIO(content))
                for name in zf.namelist():
                    if name.lower().endswith('.json') and not name.startswith('__MACOSX'):
                        json_content = zf.read(name).decode('utf-8')
                        files_data.append((os.path.basename(name), json_content))
                zf.close()
            except Exception:
                pass
        elif f.filename.lower().endswith('.json'):
            try:
                files_data.append((f.filename, content.decode('utf-8')))
            except Exception:
                pass

    # Folder path
    folder_path = request.form.get('folder_path', '').strip()

    if not files_data and not folder_path:
        return jsonify({'error': 'Chua co du lieu. Hay upload file hoac nhap duong dan thu muc.'}), 400

    jobs[job_id]['total_files'] = len(files_data)

    # Start background processing
    thread = threading.Thread(
        target=process_job,
        args=(job_id, files_data, folder_path, selected_ids, custom_patterns_list)
    )
    thread.daemon = True
    thread.start()

    return jsonify({'job_id': job_id})


@app.route('/api/status/<job_id>')
def api_status(job_id):
    """Get job status."""
    if job_id not in jobs:
        return jsonify({'error': 'Job not found'}), 404
    job = jobs[job_id]
    return jsonify({
        'status': job['status'],
        'total_files': job['total_files'],
        'processed_files': job['processed_files'],
        'current_file': job['current_file'],
        'stats': job.get('stats', {}),
    })


@app.route('/api/results/<job_id>')
def api_results(job_id):
    """Get job results (paginated)."""
    if job_id not in jobs:
        return jsonify({'error': 'Job not found'}), 404

    job = jobs[job_id]
    page = int(request.args.get('page', 1))
    per_page = int(request.args.get('per_page', 100))
    pattern_filter = request.args.get('pattern', '')

    results = job['results']
    if pattern_filter:
        results = [r for r in results if r['pattern_name'] == pattern_filter]

    total = len(results)
    start = (page - 1) * per_page
    end = start + per_page
    page_results = results[start:end]

    return jsonify({
        'total': total,
        'page': page,
        'per_page': per_page,
        'total_pages': (total + per_page - 1) // per_page,
        'results': page_results,
    })


@app.route('/api/download/<job_id>')
def api_download(job_id):
    """Download results as Excel/CSV."""
    if job_id not in jobs:
        return jsonify({'error': 'Job not found'}), 404

    job = jobs[job_id]
    if not job['results']:
        return jsonify({'error': 'No results'}), 404

    pattern_filter = request.args.get('pattern', '')
    results = job['results']
    if pattern_filter:
        results = [r for r in results if r['pattern_name'] == pattern_filter]

    output, fmt = save_to_excel(results)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

    if fmt == 'xlsx':
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'ket_qua_{timestamp}.xlsx'
        )
    else:
        return send_file(
            output,
            mimetype='text/csv',
            as_attachment=True,
            download_name=f'ket_qua_{timestamp}.csv'
        )


# ============================================================================
# MAIN
# ============================================================================

if __name__ == '__main__':
    os.makedirs('templates', exist_ok=True)
    print("=" * 60)
    print("  Korean Double Negation Extractor")
    print("  Mo trinh duyet va vao: http://localhost:5000")
    print("=" * 60)
    app.run(debug=False, host='0.0.0.0', port=5000)

